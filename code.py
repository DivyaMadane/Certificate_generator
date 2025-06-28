from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
import qrcode
from datetime import datetime
import os

def generate_qr_code_in_memory(recipient_data):
    qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=6, border=0)
    qr.add_data(recipient_data)
    qr.make(fit=True)
    qr_image = qr.make_image(fill_color="black", back_color="white")
    return qr_image

def generate_qr_code_from_data(recipient_name, event_name, certificate_date):
    if certificate_date:
        recipient_data = f"Recipient: {recipient_name} \nEvent: {event_name} \nDate: {certificate_date.strftime('%B %d, %Y')}"
        return recipient_data
    else:
        return ""

def paste_qr_code_in_memory(certificate_image, qr_image, qr_position):
    certificate_image.paste(qr_image, qr_position)

def generate_certificate(recipient_name, event_name, certificate_date):
    script_dir = os.path.dirname(os.path.abspath("background.jpg"))
    background_image_path = os.path.join(script_dir, "background.jpg")

    if os.path.exists(background_image_path):
        certificate_image = Image.open(background_image_path)
        draw = ImageDraw.Draw(certificate_image)

        bold_title_font = ImageFont.truetype("arialbd.ttf", 100)
        text_font = ImageFont.truetype("arial.ttf", 70)
        bold_text_font_path = "arial-cufonfonts\\ARIAL.TTF"
        bold_text_font = ImageFont.truetype(bold_text_font_path, 80)

        title_color = 'black'
        text_color1 = 'gold'
        text_color2 = 'black'

        title_text = "Certificate of Completion"
        line1_text = "This is to certify that"
        line2_text = "has successfully completed the event of"

        bold_title_position = (550, 300)
        line1_position = (670, 500)
        recipient_position = (700, 650)
        line2_position = (400, 800)
        event_position = (800, 925)
        line3_position = (580, 1050)

        draw.text(bold_title_position, title_text, fill=title_color, font=bold_title_font)
        draw.text(line1_position, line1_text, fill=text_color2, font=text_font)
        draw.text(recipient_position, recipient_name, fill=text_color1, font=bold_text_font)
        draw.text(line2_position, line2_text, fill=text_color2, font=text_font)
        draw.text(event_position, event_name, fill=text_color2, font=bold_text_font)

        qr_data = generate_qr_code_from_data(recipient_name, event_name, certificate_date)
        if qr_data:
            qr_image = generate_qr_code_in_memory(qr_data)
            paste_qr_code_in_memory(certificate_image, qr_image, (1520, 1000))

        draw.text(line3_position, qr_data.split("\n")[2], fill=text_color2, font=bold_text_font)

        certificate_dir = os.path.join(script_dir, "F1")
        if not os.path.exists(certificate_dir):
            os.makedirs(certificate_dir)
        certificate_image.save(os.path.join(certificate_dir, f'certificate_{recipient_name}.pdf'), "PDF", resolution=70.0)
    else:
        print("Error: Background image 'background.jpg' not found.")

def read_data_from_excel(file_path):
    try:
        if os.path.exists(file_path):
            wb = load_workbook(file_path)
            sheet = wb.active

            for row in sheet.iter_rows(min_row=2, values_only=True):
                recipient_name, event_name, certificate_date = row[:3]

                if isinstance(certificate_date, datetime):
                    generate_certificate(recipient_name, event_name, certificate_date)

        else:
            print(f"Error: Excel file '{file_path}' not found.")
    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == "__main__":
    excel_file_path = "Participant_Data.xlsx"
    read_data_from_excel(excel_file_path)
